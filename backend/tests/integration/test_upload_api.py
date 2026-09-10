from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


def workbook_bytes(code="20706100", name="测试课程", *, extra_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.cell(1, 1, "2026-2027学年第1学期")
    sheet.cell(1, 4, name + "课程课表")
    sheet.cell(1, 8, code + " 测试学院")
    for col, day in enumerate("一二三四五六日", 3):
        sheet.cell(2, col, "星期" + day)
    sheet.cell(3, 3, f"张老师/1-5周,7-8周/南望山校区 教一楼101/{name}-0001/072242/20/考试")
    sheet.cell(9, 1, "注--本学期2026-08-31正式上课至2027-01-24结束，共21周. 打印时间：2026-09-10")
    if extra_sheet:
        book.create_sheet("另一门课")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def zipped(name, data):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        z.writestr(name, data)
    return output.getvalue()


def test_xlsx_direct_import_enters_planning_payload(client, auth_headers):
    response = client.post(
        "/api/catalog/import-courses",
        headers=auth_headers,
        files=[("files", ("测试.xlsx", workbook_bytes()))],
    )
    assert response.status_code == 200, response.text
    choices = response.json()["courses"]
    meeting = choices[0]["custom"]["sections"][0]["meetings"][0]
    assert meeting["weeks"] == [1, 2, 3, 4, 5, 7, 8]
    client.put(
        "/api/profile",
        headers=auth_headers,
        json={"college": "工程学院", "major": "土木工程", "cohort_year": 2024},
    )
    plan = client.post(
        "/api/plans/generate",
        headers=auth_headers,
        json={"input_mode": "manual", "manual_courses": choices},
    )
    assert plan.status_code == 200, plan.text
    assert plan.json()["plans"][0]["scheduled_course_count"] == 1


def test_multiple_zips_are_one_snapshot_and_failure_is_atomic(client, auth_headers):
    a = zipped("a.xlsx", workbook_bytes())
    b = zipped("b.xlsx", workbook_bytes("20706101", "另一课程"))
    response = client.post(
        "/api/catalog/import",
        headers=auth_headers,
        files=[("files", ("one.zip", a)), ("files", ("two.zip", b))],
    )
    assert response.status_code == 200, response.text
    assert response.json()["courses"] == 2
    assert response.json()["snapshots"] == 1
    assert response.json()["default_eligible_sections"] == 2
    failed = client.post(
        "/api/catalog/import",
        headers=auth_headers,
        files=[("files", ("one.zip", a)), ("files", ("bad.zip", b"bad"))],
    )
    assert failed.status_code == 422
    assert client.get("/api/catalog/status", headers=auth_headers).json()["course_count"] == 2


def test_multisheet_is_rejected_instead_of_silently_ignoring_data(client, auth_headers):
    result = client.post(
        "/api/catalog/import-courses",
        headers=auth_headers,
        files=[("files", ("multi.xlsx", workbook_bytes(extra_sheet=True)))],
    )
    assert result.status_code == 422
    assert "一个工作表" in result.json()["detail"]


def test_unsafe_zip_member_rejects_the_whole_batch(client, auth_headers):
    result = client.post(
        "/api/catalog/import",
        headers=auth_headers,
        files=[("files", ("bad.zip", zipped("../a.xlsx", workbook_bytes())))],
    )
    assert result.status_code == 422


def test_downloaded_standard_template_roundtrips_without_teacher_or_location(client, auth_headers):
    response = client.get("/api/catalog/template", headers=auth_headers)
    assert response.status_code == 200
    imported = client.post(
        "/api/catalog/import-courses",
        headers=auth_headers,
        files=[("files", ("template.xlsx", response.content))],
    )
    assert imported.status_code == 200, imported.text
    section = imported.json()["courses"][0]["custom"]["sections"][0]
    assert section["instructors"] == []
    assert len(section["meetings"]) == 2
    assert section["meetings"][1]["weeks"] == [1, 2, 3, 6, 7, 8]


def test_pdf_file_is_rejected_without_loading_a_pdf_parser(client, auth_headers):
    response = client.post(
        "/api/catalog/import-courses",
        headers=auth_headers,
        files=[("files", ("schedule.pdf", b"%PDF-1.7"))],
    )
    assert response.status_code == 422


def test_suffix_components_require_confirmation_after_import(client, auth_headers):
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(workbook_bytes()))
    book.active.cell(3, 3, "张老师/1-5周/南望山校区 教一楼101/测试课程-0001A/072242/20/考试")
    content = BytesIO()
    book.save(content)
    imported = client.post(
        "/api/catalog/import-courses",
        headers=auth_headers,
        files=[("files", ("a.xlsx", content.getvalue()))],
    )
    assert imported.status_code == 200
    assert imported.json()["courses"][0]["custom"]["component_relationship_confirmed"] is False
